from __future__ import annotations

from datetime import UTC, datetime
from decimal import Decimal

from openpyxl import load_workbook
from sqlalchemy import delete, select

from germania.collectors.autoscout24 import AutoScout24ListingImportService
from germania.collectors.marketplace import MarketplaceListingRecord
from germania.db import (
    Base,
    MarketplaceListing,
    MarketplacePriceHistory,
    create_database_engine,
    create_session_factory,
    session_scope,
)
from germania.db.seed import seed_configuration
from germania.quality.marketplace import create_marketplace_quality_report


def test_quality_report_detects_anomalies_and_relationship_mismatch(tmp_path) -> None:
    engine = create_database_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    session_factory = create_session_factory(engine)
    collected_at = datetime(2026, 7, 19, 8, tzinfo=UTC)

    try:
        with session_scope(session_factory) as session:
            seed_configuration(session)
            imported = AutoScout24ListingImportService(session).import_records(
                [
                    _record("quality-1", "500", 600001, collected_at),
                    _record("quality-2", "30000", None, collected_at),
                ]
            )
            assert imported.inserted == 2
            second = session.scalar(
                select(MarketplaceListing).where(
                    MarketplaceListing.external_listing_id == "quality-2"
                )
            )
            assert second is not None
            second.model_name = "Wrong Model"
            session.execute(
                delete(MarketplacePriceHistory).where(
                    MarketplacePriceHistory.marketplace_listing_id
                    == second.marketplace_listing_id
                )
            )
            result = create_marketplace_quality_report(session, tmp_path)

        assert result.total_listings == 2
        assert result.overall_completeness_rate < 1
        assert result.duplicate_external_id_groups == 0
        assert result.listings_without_price_history == 1
        assert result.price_anomalies == 1
        assert result.mileage_anomalies == 1
        assert result.brand_model_inconsistencies == 1

        workbook = load_workbook(result.workbook_path, read_only=False, data_only=True)
        try:
            assert workbook.sheetnames == [
                "Quality_Summary",
                "Field_Completeness",
                "Missing_Values",
                "Duplicate_IDs",
                "Missing_History",
                "Price_Anomalies",
                "Mileage_Anomalies",
                "Brand_Model_Check",
            ]
            assert workbook["Price_Anomalies"].max_row == 2
            assert workbook["Mileage_Anomalies"].max_row == 2
            assert workbook["Brand_Model_Check"].max_row == 2
        finally:
            workbook.close()
    finally:
        Base.metadata.drop_all(engine)
        engine.dispose()


def _record(
    listing_id: str,
    price: str,
    mileage: int | None,
    collected_at: datetime,
) -> MarketplaceListingRecord:
    return MarketplaceListingRecord(
        source_id="autoscout24_de",
        external_listing_id=listing_id,
        collected_at=collected_at,
        listing_url=None,
        brand_name="Volkswagen",
        model_name="Golf",
        price_amount=Decimal(price),
        mileage_km=mileage,
    )
