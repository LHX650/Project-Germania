from __future__ import annotations

from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook

from germania.collectors.kba import (
    KBAXlsxParser,
    KBAXlsxParserError,
    RegistrationRecord,
    parse_kba_fz10_xlsx,
)

FIXTURE_DIR = Path(__file__).resolve().parents[3] / "fixtures" / "kba"
FZ10_FIXTURE = FIXTURE_DIR / "fz10_2026_06_sample.xlsx"
COLLECTED_AT = datetime(2026, 7, 18, 9, 15, tzinfo=UTC)
SOURCE_URL = "https://www.kba.de/monthly/fz10_2026_06.xlsx"


def test_parse_kba_fz10_xlsx_fixture_outputs_registration_records() -> None:
    records = parse_kba_fz10_xlsx(
        FZ10_FIXTURE,
        source_url=SOURCE_URL,
        collected_at=COLLECTED_AT,
    )

    assert len(records) == 5
    assert all(isinstance(record, RegistrationRecord) for record in records)
    assert all(record.source_id == "kba" for record in records)
    assert all(record.year == 2026 for record in records)
    assert all(record.month == 6 for record in records)
    assert all(record.collected_at == COLLECTED_AT for record in records)


def test_parser_auto_detects_header_and_long_format_fields() -> None:
    records = parse_kba_fz10_xlsx(
        FZ10_FIXTURE,
        source_url=SOURCE_URL,
        collected_at=COLLECTED_AT,
    )

    first = records[0]
    assert first.brand == "Volkswagen"
    assert first.model_series == "Golf"
    assert first.registrations == 1234
    assert first.fuel_type == "total"
    assert first.market_share == Decimal("4.5")
    assert first.source_url == "https://www.kba.de/fz10-test.xlsx"

    electric_golf = records[1]
    assert electric_golf.brand == "Volkswagen"
    assert electric_golf.model_series == "Golf"
    assert electric_golf.registrations == 120
    assert electric_golf.fuel_type == "battery_electric"
    assert electric_golf.market_share == Decimal("0.4")
    assert electric_golf.source_url == SOURCE_URL


def test_parser_expands_wide_fuel_columns() -> None:
    records = parse_kba_fz10_xlsx(
        FZ10_FIXTURE,
        source_url=SOURCE_URL,
        collected_at=COLLECTED_AT,
    )
    enyaq_records = [
        record
        for record in records
        if record.brand == "Skoda" and record.model_series == "Enyaq"
    ]

    assert len(enyaq_records) == 2
    assert {record.fuel_type for record in enyaq_records} == {
        "total",
        "battery_electric",
    }
    assert {record.registrations for record in enyaq_records} == {77}
    assert {record.market_share for record in enyaq_records} == {Decimal("0.3")}


def test_parser_skips_total_rows_and_footnotes() -> None:
    records = parse_kba_fz10_xlsx(
        FZ10_FIXTURE,
        source_url=SOURCE_URL,
        collected_at=COLLECTED_AT,
    )

    assert all(record.brand not in {"Insgesamt", "Gesamt"} for record in records)
    assert all("Fußnote" not in record.brand for record in records)


def test_parser_safely_handles_missing_optional_values(tmp_path: Path) -> None:
    workbook_path = tmp_path / "fz10_2026_07_missing_values.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Juli 2026"
    worksheet.append(["Berichtsmonat: Juli 2026"])
    worksheet.append(["Hersteller", "Handelsname", "Neuzulassungen"])
    worksheet.append(["Volkswagen", "Tiguan", "2.001"])
    workbook.save(workbook_path)

    records = KBAXlsxParser().parse_file(workbook_path, collected_at=COLLECTED_AT)

    assert len(records) == 1
    record = records[0]
    assert record.brand == "Volkswagen"
    assert record.model_series == "Tiguan"
    assert record.registrations == 2001
    assert record.fuel_type == "total"
    assert record.market_share is None
    assert record.source_url is None


def test_parser_returns_empty_list_for_workbook_without_fz10_table(
    tmp_path: Path,
) -> None:
    workbook_path = tmp_path / "fz10_2026_08_empty.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "No table"
    worksheet.append(["Berichtsmonat: August 2026"])
    worksheet.append(["Only a note, no parseable KBA table"])
    workbook.save(workbook_path)

    records = parse_kba_fz10_xlsx(workbook_path, collected_at=COLLECTED_AT)

    assert records == []


def test_parser_rejects_missing_file(tmp_path: Path) -> None:
    with pytest.raises(KBAXlsxParserError, match="not found"):
        parse_kba_fz10_xlsx(tmp_path / "missing.xlsx")


def test_parser_requires_reporting_period_when_not_in_workbook(
    tmp_path: Path,
) -> None:
    workbook_path = tmp_path / "fz10_no_period.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "FZ10"
    worksheet.append(["Hersteller", "Handelsname", "Neuzulassungen"])
    worksheet.append(["Volkswagen", "Golf", 100])
    workbook.save(workbook_path)

    with pytest.raises(KBAXlsxParserError, match="reporting period"):
        parse_kba_fz10_xlsx(workbook_path)


def test_parser_accepts_explicit_reporting_period(tmp_path: Path) -> None:
    workbook_path = tmp_path / "fz10_no_period.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "FZ10"
    worksheet.append(["Hersteller", "Handelsname", "Neuzulassungen"])
    worksheet.append(["Volkswagen", "Golf", 100])
    workbook.save(workbook_path)

    records = parse_kba_fz10_xlsx(
        workbook_path,
        year=2026,
        month=9,
        collected_at=COLLECTED_AT,
    )

    assert len(records) == 1
    assert records[0].year == 2026
    assert records[0].month == 9
