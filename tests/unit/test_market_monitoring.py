from __future__ import annotations

from datetime import UTC, datetime
from decimal import Decimal

from openpyxl import load_workbook
from sqlalchemy import func, select

from germania.collectors.autoscout24 import AutoScout24ListingImportService
from germania.collectors.marketplace import MarketplaceListingRecord
from germania.db import (
    Base,
    MarketplaceListing,
    MarketplacePriceHistory,
    create_database_engine,
    create_session_factory,
    session_scope,
)
from germania.db.seed import seed_configuration
from germania.market_monitoring import create_market_monitor_report


def test_market_monitor_compares_baseline_and_current_window(tmp_path) -> None:
    engine = create_database_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    session_factory = create_session_factory(engine)
    baseline_at = datetime(2026, 7, 19, 8, tzinfo=UTC)
    cutoff = datetime(2026, 7, 19, 9, tzinfo=UTC)
    current_at = datetime(2026, 7, 19, 10, tzinfo=UTC)

    try:
        with session_scope(session_factory) as session:
            seed_configuration(session)
            service = AutoScout24ListingImportService(session)
            baseline = service.import_records(
                [
                    _record("golf-stay", "Volkswagen", "Golf", "20000", baseline_at),
                    _record("golf-removed", "Volkswagen", "Golf", "15000", baseline_at),
                    _record("model-y-stay", "Tesla", "Model Y", "40000", baseline_at),
                ]
            )
            assert baseline.inserted == 3
            current = service.import_records(
                [
                    _record("golf-stay", "Volkswagen", "Golf", "18000", current_at),
                    _record("golf-new", "Volkswagen", "Golf", "22000", current_at),
                    _record("model-y-stay", "Tesla", "Model Y", "45000", current_at),
                ]
            )
            assert current.inserted == 1
            assert current.updated == 2
            counts_before = _database_counts(session)

            result = create_market_monitor_report(
                session,
                cutoff=cutoff,
                current_end=current_at,
                output_dir=tmp_path,
            )

            assert _database_counts(session) == counts_before

        assert result.previous_listings == 3
        assert result.current_listings == 3
        assert result.new_listings == 1
        assert result.removed_listings == 1
        assert result.price_decreases == 1
        assert result.price_increases == 1
        assert result.unchanged_prices == 0
        assert result.total_price_change == Decimal("3000.00")
        assert result.average_price_change == Decimal("1500.00")

        workbook = load_workbook(result.workbook_path, read_only=False, data_only=True)
        try:
            assert workbook.sheetnames == [
                "Summary",
                "New_Listings",
                "Removed_Listings",
                "Price_Changes",
                "Brand_Changes",
                "Vehicle_Changes",
            ]
            assert workbook["New_Listings"].max_row == 2
            assert workbook["Removed_Listings"].max_row == 2
            assert workbook["Price_Changes"].max_row == 3
            assert workbook["Brand_Changes"].max_row == 3
            assert workbook["Vehicle_Changes"].max_row == 3
        finally:
            workbook.close()
    finally:
        Base.metadata.drop_all(engine)
        engine.dispose()


def _record(
    listing_id: str,
    brand: str,
    model: str,
    price: str,
    collected_at: datetime,
) -> MarketplaceListingRecord:
    return MarketplaceListingRecord(
        source_id="autoscout24_de",
        external_listing_id=listing_id,
        collected_at=collected_at,
        listing_url=f"https://www.autoscout24.de/angebote/{listing_id}",
        brand_name=brand,
        model_name=model,
        variant_name="Test variant",
        title=f"{brand} {model} Test variant",
        price_amount=Decimal(price),
        mileage_km=10000,
    )


def _database_counts(session) -> tuple[int, int]:
    return (
        int(session.scalar(select(func.count()).select_from(MarketplaceListing)) or 0),
        int(
            session.scalar(select(func.count()).select_from(MarketplacePriceHistory))
            or 0
        ),
    )
