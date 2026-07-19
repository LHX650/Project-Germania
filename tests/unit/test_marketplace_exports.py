from __future__ import annotations

import csv
from datetime import UTC, datetime
from decimal import Decimal

from openpyxl import load_workbook

from germania.collectors.autoscout24 import AutoScout24ListingImportService
from germania.collectors.marketplace import MarketplaceListingRecord
from germania.db import (
    Base,
    create_database_engine,
    create_session_factory,
    session_scope,
)
from germania.db.seed import seed_configuration
from germania.marketplace_exports import export_marketplace_data


def test_marketplace_export_writes_expected_excel_and_csv(tmp_path) -> None:
    engine = create_database_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    session_factory = create_session_factory(engine)
    collected_at = datetime(2026, 7, 19, 8, tzinfo=UTC)

    try:
        with session_scope(session_factory) as session:
            seed_configuration(session)
            imported = AutoScout24ListingImportService(session).import_records(
                [
                    _record("golf-export-1", "20000", 10000, collected_at),
                    _record("golf-export-2", "30000", 20000, collected_at),
                ]
            )
            assert imported.inserted == 2
            result = export_marketplace_data(session, tmp_path)

        assert result.listing_count == 2
        assert result.price_history_count == 2
        assert {path.name for path in result.csv_paths} == {
            "listings.csv",
            "price_history.csv",
            "brand_summary.csv",
            "vehicle_summary.csv",
        }
        with (tmp_path / "listings.csv").open(
            encoding="utf-8-sig", newline=""
        ) as csv_file:
            rows = list(csv.reader(csv_file))
        assert len(rows) == 3
        assert rows[0][2] == "external_listing_id"

        workbook = load_workbook(result.workbook_path, read_only=False, data_only=True)
        try:
            assert workbook.sheetnames == [
                "Listings",
                "Price_History",
                "Brand_Summary",
                "Vehicle_Summary",
                "Collection_Summary",
            ]
            assert workbook["Listings"].max_row == 3
            assert workbook["Price_History"].max_row == 3
            collection_values = {
                row[0]: row[1]
                for row in workbook["Collection_Summary"].iter_rows(
                    min_row=2, values_only=True
                )
            }
            assert collection_values["marketplace_listings"] == 2
            assert collection_values["price_history"] == 2
        finally:
            workbook.close()
    finally:
        Base.metadata.drop_all(engine)
        engine.dispose()


def _record(
    listing_id: str,
    price: str,
    mileage: int,
    collected_at: datetime,
) -> MarketplaceListingRecord:
    return MarketplaceListingRecord(
        source_id="autoscout24_de",
        external_listing_id=listing_id,
        collected_at=collected_at,
        listing_url=f"https://www.autoscout24.de/angebote/{listing_id}",
        brand_name="Volkswagen",
        model_name="Golf",
        variant_name="1.5 TSI",
        title="Volkswagen Golf 1.5 TSI",
        price_amount=Decimal(price),
        registration_year=2022,
        mileage_km=mileage,
    )
