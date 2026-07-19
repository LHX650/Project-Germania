"""Read-only comparison of two marketplace collection result windows."""

from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl.styles import Alignment
from sqlalchemy import select
from sqlalchemy.orm import Session

from germania.db.models import MarketplaceListing, MarketplacePriceHistory
from germania.utils.workbook import (
    add_table_sheet,
    create_report_workbook,
    save_report_workbook,
)

SUMMARY_HEADERS = ("metric", "value", "definition")
LISTING_HEADERS = (
    "marketplace_listing_id",
    "external_listing_id",
    "brand_name",
    "model_name",
    "variant_name",
    "price_amount",
    "mileage_km",
    "first_seen_at",
    "last_seen_at",
    "listing_url",
    "status_basis",
)
PRICE_CHANGE_HEADERS = (
    "marketplace_listing_id",
    "external_listing_id",
    "brand_name",
    "model_name",
    "variant_name",
    "previous_price",
    "current_price",
    "price_change_amount",
    "absolute_price_change",
    "price_change_percent",
    "direction",
    "listing_url",
)
BRAND_CHANGE_HEADERS = (
    "brand_name",
    "previous_listings",
    "current_listings",
    "new_listings",
    "removed_listings",
    "net_listing_change",
    "previous_average_price",
    "current_average_price",
    "average_market_price_change",
    "average_matched_listing_price_change",
)
VEHICLE_CHANGE_HEADERS = (
    "brand_name",
    "model_name",
    *BRAND_CHANGE_HEADERS[1:],
)


@dataclass(frozen=True)
class MarketMonitorResult:
    """Summary and report path for one market comparison."""

    workbook_path: Path
    cutoff: datetime
    current_end: datetime
    scoped_brands: int
    scoped_vehicles: int
    previous_listings: int
    current_listings: int
    new_listings: int
    removed_listings: int
    price_decreases: int
    price_increases: int
    unchanged_prices: int
    total_price_change: Decimal
    average_price_change: Decimal | None


def create_market_monitor_report(
    session: Session,
    *,
    cutoff: datetime,
    current_end: datetime | None = None,
    output_dir: Path | str = Path("exports"),
) -> MarketMonitorResult:
    """Compare a baseline cohort with listings observed after ``cutoff``.

    The comparison scope is restricted to brand/model pairs actually observed
    in the current window. A baseline listing not observed in that window is
    reported as inferred removed, not as a confirmed sale or permanent delist.
    """

    boundary = _normalize_datetime(cutoff, "cutoff")
    end = _normalize_datetime(current_end, "current_end") if current_end else None
    if end is not None and end <= boundary:
        raise ValueError("current_end must be later than cutoff")

    latest = session.scalar(
        select(MarketplaceListing.last_collected_at)
        .order_by(MarketplaceListing.last_collected_at.desc())
        .limit(1)
    )
    if latest is None:
        raise ValueError("marketplace database contains no collected listings")
    comparison_end = end or _normalize_datetime(latest, "latest collection time")
    if comparison_end <= boundary:
        raise ValueError("cutoff must be earlier than the latest collection time")

    current_candidates = list(
        session.scalars(
            select(MarketplaceListing).where(
                MarketplaceListing.last_collected_at >= boundary,
                MarketplaceListing.last_collected_at <= comparison_end,
            )
        )
    )
    scope = {
        (listing.brand_name, listing.model_name)
        for listing in current_candidates
        if listing.brand_name and listing.model_name
    }
    if not scope:
        raise ValueError("current window contains no comparable brand/model pairs")

    all_listings = list(
        session.scalars(
            select(MarketplaceListing).order_by(
                MarketplaceListing.marketplace_listing_id
            )
        )
    )
    scoped_listings = [
        listing
        for listing in all_listings
        if (listing.brand_name, listing.model_name) in scope
    ]
    previous = {
        listing.marketplace_listing_id: listing
        for listing in scoped_listings
        if listing.first_seen_at is not None
        and _normalize_datetime(listing.first_seen_at, "first_seen_at") < boundary
    }
    current = {
        listing.marketplace_listing_id: listing
        for listing in scoped_listings
        if listing.last_collected_at is not None
        and boundary
        <= _normalize_datetime(listing.last_collected_at, "last_collected_at")
        <= comparison_end
    }
    previous_prices = _previous_prices(session, tuple(previous), boundary)

    previous_ids = set(previous)
    current_ids = set(current)
    new_ids = current_ids - previous_ids
    removed_ids = previous_ids - current_ids
    shared_ids = previous_ids & current_ids
    price_changes = _price_change_rows(
        shared_ids,
        current,
        previous_prices,
    )
    decrease_count = sum(row[10] == "decrease" for row in price_changes)
    increase_count = sum(row[10] == "increase" for row in price_changes)
    comparable_shared = {
        listing_id
        for listing_id in shared_ids
        if listing_id in previous_prices
        and current[listing_id].current_price_amount is not None
    }
    unchanged_count = len(comparable_shared) - len(price_changes)
    signed_changes = [row[7] for row in price_changes]
    total_change = sum(signed_changes, start=Decimal("0"))
    average_change = (
        total_change / len(signed_changes) if signed_changes else Decimal("0")
    )

    new_rows = [
        _listing_row(current[item], "first seen in current window")
        for item in sorted(new_ids)
    ]
    removed_rows = [
        _listing_row(
            previous[item],
            "not observed in current window (inferred)",
            previous_prices.get(item),
        )
        for item in sorted(removed_ids)
    ]
    brand_rows = _group_change_rows(
        previous,
        current,
        previous_prices,
        new_ids,
        removed_ids,
        key=lambda listing: (listing.brand_name or "Unknown",),
    )
    vehicle_rows = _group_change_rows(
        previous,
        current,
        previous_prices,
        new_ids,
        removed_ids,
        key=lambda listing: (
            listing.brand_name or "Unknown",
            listing.model_name or "Unknown",
        ),
    )
    summary_rows = [
        ("cutoff", boundary, "Baseline observations are before this UTC timestamp."),
        (
            "current_end",
            comparison_end,
            "Current comparison window ends at this UTC timestamp.",
        ),
        (
            "scoped_brands",
            len({item[0] for item in scope}),
            "Brands observed in the current window.",
        ),
        (
            "scoped_vehicles",
            len(scope),
            "Brand/model pairs observed in the current window.",
        ),
        ("previous_listings", len(previous), "Baseline cohort in the compared scope."),
        ("current_listings", len(current), "Listings observed in the current window."),
        ("new_listings", len(new_ids), "Current IDs absent from the baseline cohort."),
        (
            "removed_listings",
            len(removed_ids),
            "Baseline IDs not observed in the current window; inferred only.",
        ),
        (
            "net_listing_change",
            len(current) - len(previous),
            "Current listings minus baseline listings.",
        ),
        (
            "price_decreases",
            decrease_count,
            "Shared listings with a lower current asking price.",
        ),
        (
            "price_increases",
            increase_count,
            "Shared listings with a higher current asking price.",
        ),
        (
            "unchanged_prices",
            unchanged_count,
            "Comparable shared listings with unchanged asking price.",
        ),
        (
            "total_price_change",
            total_change,
            "Signed sum across changed shared listings.",
        ),
        (
            "average_price_change",
            average_change,
            "Average signed change across changed shared listings.",
        ),
        (
            "method",
            "first_seen/last_collected cohort comparison",
            "No CollectionBatch or ListingObservation rows exist in this database.",
        ),
    ]

    workbook = create_report_workbook()
    summary_sheet = add_table_sheet(
        workbook,
        title="Summary",
        headers=SUMMARY_HEADERS,
        rows=summary_rows,
        table_name="MonitorSummaryTable",
    )
    summary_sheet.column_dimensions["B"].width = 38
    summary_sheet.column_dimensions["C"].width = 64
    for row in range(2, summary_sheet.max_row + 1):
        summary_sheet.cell(row=row, column=2).alignment = Alignment(
            wrap_text=True,
            vertical="top",
        )
        summary_sheet.cell(row=row, column=3).alignment = Alignment(
            wrap_text=True,
            vertical="top",
        )
        summary_sheet.row_dimensions[row].height = 30
    add_table_sheet(
        workbook,
        title="New_Listings",
        headers=LISTING_HEADERS,
        rows=new_rows,
        table_name="NewListingsTable",
    )
    add_table_sheet(
        workbook,
        title="Removed_Listings",
        headers=LISTING_HEADERS,
        rows=removed_rows,
        table_name="RemovedListingsTable",
    )
    add_table_sheet(
        workbook,
        title="Price_Changes",
        headers=PRICE_CHANGE_HEADERS,
        rows=price_changes,
        table_name="PriceChangesTable",
    )
    add_table_sheet(
        workbook,
        title="Brand_Changes",
        headers=BRAND_CHANGE_HEADERS,
        rows=brand_rows,
        table_name="BrandChangesTable",
    )
    add_table_sheet(
        workbook,
        title="Vehicle_Changes",
        headers=VEHICLE_CHANGE_HEADERS,
        rows=vehicle_rows,
        table_name="VehicleChangesTable",
    )
    workbook_path = Path(output_dir) / "market_monitor_report.xlsx"
    save_report_workbook(workbook, workbook_path)

    return MarketMonitorResult(
        workbook_path=workbook_path,
        cutoff=boundary,
        current_end=comparison_end,
        scoped_brands=len({item[0] for item in scope}),
        scoped_vehicles=len(scope),
        previous_listings=len(previous),
        current_listings=len(current),
        new_listings=len(new_ids),
        removed_listings=len(removed_ids),
        price_decreases=decrease_count,
        price_increases=increase_count,
        unchanged_prices=unchanged_count,
        total_price_change=total_change,
        average_price_change=average_change,
    )


def _previous_prices(
    session: Session,
    listing_ids: tuple[int, ...],
    cutoff: datetime,
) -> dict[int, Decimal]:
    if not listing_ids:
        return {}
    rows = session.execute(
        select(MarketplacePriceHistory)
        .where(
            MarketplacePriceHistory.marketplace_listing_id.in_(listing_ids),
            MarketplacePriceHistory.observed_at < cutoff,
        )
        .order_by(
            MarketplacePriceHistory.marketplace_listing_id,
            MarketplacePriceHistory.observed_at,
        )
    ).scalars()
    output: dict[int, Decimal] = {}
    for row in rows:
        output[row.marketplace_listing_id] = row.price_amount
    return output


def _price_change_rows(
    shared_ids: set[int],
    current: dict[int, MarketplaceListing],
    previous_prices: dict[int, Decimal],
) -> list[tuple[Any, ...]]:
    output: list[tuple[Any, ...]] = []
    for listing_id in sorted(shared_ids):
        listing = current[listing_id]
        previous_price = previous_prices.get(listing_id)
        current_price = listing.current_price_amount
        if (
            previous_price is None
            or current_price is None
            or previous_price == current_price
        ):
            continue
        change = current_price - previous_price
        output.append(
            (
                listing.marketplace_listing_id,
                listing.external_listing_id,
                listing.brand_name,
                listing.model_name,
                listing.variant_name,
                previous_price,
                current_price,
                change,
                abs(change),
                change / previous_price if previous_price else None,
                "decrease" if change < 0 else "increase",
                listing.listing_url,
            )
        )
    return output


def _listing_row(
    listing: MarketplaceListing,
    status_basis: str,
    price: Decimal | None = None,
) -> tuple[Any, ...]:
    return (
        listing.marketplace_listing_id,
        listing.external_listing_id,
        listing.brand_name,
        listing.model_name,
        listing.variant_name,
        price if price is not None else listing.current_price_amount,
        listing.mileage_km,
        listing.first_seen_at,
        listing.last_seen_at,
        listing.listing_url,
        status_basis,
    )


def _group_change_rows(
    previous: dict[int, MarketplaceListing],
    current: dict[int, MarketplaceListing],
    previous_prices: dict[int, Decimal],
    new_ids: set[int],
    removed_ids: set[int],
    *,
    key: Any,
) -> list[tuple[Any, ...]]:
    previous_groups: dict[tuple[str, ...], list[int]] = defaultdict(list)
    current_groups: dict[tuple[str, ...], list[int]] = defaultdict(list)
    for listing_id, listing in previous.items():
        previous_groups[key(listing)].append(listing_id)
    for listing_id, listing in current.items():
        current_groups[key(listing)].append(listing_id)

    output: list[tuple[Any, ...]] = []
    for group in sorted(set(previous_groups) | set(current_groups)):
        previous_ids = previous_groups[group]
        current_ids = current_groups[group]
        previous_values = [
            previous_prices[item] for item in previous_ids if item in previous_prices
        ]
        current_values = [
            current[item].current_price_amount
            for item in current_ids
            if current[item].current_price_amount is not None
        ]
        shared_values = [
            current[item].current_price_amount - previous_prices[item]
            for item in set(previous_ids) & set(current_ids)
            if item in previous_prices
            and current[item].current_price_amount is not None
        ]
        previous_average = _average(previous_values)
        current_average = _average(current_values)
        output.append(
            (
                *group,
                len(previous_ids),
                len(current_ids),
                len(set(current_ids) & new_ids),
                len(set(previous_ids) & removed_ids),
                len(current_ids) - len(previous_ids),
                previous_average,
                current_average,
                (
                    current_average - previous_average
                    if previous_average is not None and current_average is not None
                    else None
                ),
                _average(shared_values),
            )
        )
    return output


def _average(values: list[Decimal]) -> Decimal | None:
    return sum(values, start=Decimal("0")) / len(values) if values else None


def _normalize_datetime(value: datetime, field: str) -> datetime:
    if not isinstance(value, datetime):
        raise ValueError(f"{field} must be a datetime")
    if value.tzinfo is not None:
        return value.astimezone(UTC).replace(tzinfo=None)
    return value
