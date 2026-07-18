"""Parser for local KBA FZ10 monthly registration XLSX files."""

from __future__ import annotations

import logging
import re
from collections.abc import Sequence
from datetime import UTC, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from germania.collectors.kba.field_mapping import (
    DEFAULT_FUEL_TYPE,
    FOOTNOTE_MARKERS,
    GERMAN_MONTHS,
    KBA_SOURCE_ID,
    SKIP_ROW_MARKERS,
    HeaderMapping,
    normalize_header,
    resolve_field_name,
    resolve_fuel_column,
    resolve_fuel_type_value,
)
from germania.collectors.kba.models import RegistrationRecord

logger = logging.getLogger(__name__)

_PERIOD_NUMERIC_PATTERN = re.compile(
    r"(?P<year>20\d{2})[-_ ./](?P<month>0?[1-9]|1[0-2])"
)
_YEAR_PATTERN = re.compile(r"\b(20\d{2})\b")
_NUMBER_PATTERN = re.compile(r"-?\d[\d\s.,]*")


class KBAXlsxParserError(ValueError):
    """Raised when a KBA workbook cannot be parsed safely."""


class KBAXlsxParser:
    """Parse KBA FZ10 monthly registration workbooks from local XLSX files."""

    def parse_file(
        self,
        path: Path | str,
        *,
        year: int | None = None,
        month: int | None = None,
        source_url: str | None = None,
        collected_at: datetime | None = None,
    ) -> list[RegistrationRecord]:
        """Parse a local KBA FZ10 XLSX file into RegistrationRecord objects."""
        workbook_path = Path(path)
        if not workbook_path.exists():
            raise KBAXlsxParserError(f"KBA workbook not found: {workbook_path}")

        parsed_at = _resolve_collected_at(collected_at)
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        try:
            records: list[RegistrationRecord] = []
            for sheet in workbook.worksheets:
                rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
                if not rows:
                    continue
                period = _resolve_period(rows, workbook_path, sheet.title, year, month)
                header_mapping = _find_header_mapping(rows)
                if header_mapping is None:
                    logger.debug("No KBA FZ10 header found in sheet=%s", sheet.title)
                    continue
                records.extend(
                    _records_from_rows(
                        rows,
                        header_mapping,
                        year=period[0],
                        month=period[1],
                        source_url=source_url,
                        collected_at=parsed_at,
                    )
                )
            return records
        finally:
            workbook.close()


def parse_kba_fz10_xlsx(
    path: Path | str,
    *,
    year: int | None = None,
    month: int | None = None,
    source_url: str | None = None,
    collected_at: datetime | None = None,
) -> list[RegistrationRecord]:
    """Parse a local KBA FZ10 XLSX file into RegistrationRecord objects."""
    return KBAXlsxParser().parse_file(
        path,
        year=year,
        month=month,
        source_url=source_url,
        collected_at=collected_at,
    )


def _find_header_mapping(rows: Sequence[Sequence[object]]) -> HeaderMapping | None:
    for row_index, row in enumerate(rows):
        columns: dict[str, int] = {}
        fuel_value_columns: dict[str, int] = {}

        for column_index, value in enumerate(row):
            field_name = resolve_field_name(value)
            if field_name is not None and field_name not in columns:
                columns[field_name] = column_index

            fuel_type = resolve_fuel_column(value)
            if fuel_type is not None:
                fuel_value_columns[fuel_type] = column_index

        has_identity = {"brand", "model_series"}.issubset(columns)
        has_registration_value = "registrations" in columns or bool(fuel_value_columns)
        if has_identity and has_registration_value:
            return HeaderMapping(
                header_row_index=row_index,
                columns=columns,
                fuel_value_columns=fuel_value_columns,
            )

    return None


def _records_from_rows(
    rows: Sequence[Sequence[object]],
    header_mapping: HeaderMapping,
    *,
    year: int,
    month: int,
    source_url: str | None,
    collected_at: datetime,
) -> list[RegistrationRecord]:
    records: list[RegistrationRecord] = []

    for row in rows[header_mapping.header_row_index + 1 :]:
        if _is_blank_or_footnote_row(row):
            continue
        brand = _cell_text(row, header_mapping.columns.get("brand"))
        model_series = _cell_text(row, header_mapping.columns.get("model_series"))
        if _should_skip_data_row(brand, model_series):
            continue

        if (
            "fuel_type" in header_mapping.columns
            or not header_mapping.fuel_value_columns
        ):
            record = _record_from_long_row(
                row,
                header_mapping,
                year=year,
                month=month,
                brand=brand,
                model_series=model_series,
                source_url=source_url,
                collected_at=collected_at,
            )
            if record is not None:
                records.append(record)
            continue

        records.extend(
            _records_from_wide_row(
                row,
                header_mapping,
                year=year,
                month=month,
                brand=brand,
                model_series=model_series,
                source_url=source_url,
                collected_at=collected_at,
            )
        )

    return records


def _record_from_long_row(
    row: Sequence[object],
    header_mapping: HeaderMapping,
    *,
    year: int,
    month: int,
    brand: str,
    model_series: str,
    source_url: str | None,
    collected_at: datetime,
) -> RegistrationRecord | None:
    registrations = _parse_integer(
        _cell_value(row, header_mapping.columns["registrations"])
    )
    if registrations is None:
        return None

    source_url_from_row = _cell_text(row, header_mapping.columns.get("source_url"))
    return RegistrationRecord(
        source_id=KBA_SOURCE_ID,
        year=_parse_year(_cell_value(row, header_mapping.columns.get("year"))) or year,
        month=_parse_month(_cell_value(row, header_mapping.columns.get("month")))
        or month,
        brand=brand,
        model_series=model_series,
        registrations=registrations,
        fuel_type=resolve_fuel_type_value(
            _cell_value(row, header_mapping.columns.get("fuel_type"))
        )
        or DEFAULT_FUEL_TYPE,
        market_share=_parse_decimal_percent(
            _cell_value(row, header_mapping.columns.get("market_share"))
        ),
        source_url=source_url_from_row or source_url,
        collected_at=collected_at,
    )


def _records_from_wide_row(
    row: Sequence[object],
    header_mapping: HeaderMapping,
    *,
    year: int,
    month: int,
    brand: str,
    model_series: str,
    source_url: str | None,
    collected_at: datetime,
) -> list[RegistrationRecord]:
    records: list[RegistrationRecord] = []
    source_url_from_row = _cell_text(row, header_mapping.columns.get("source_url"))

    for fuel_type, column_index in header_mapping.fuel_value_columns.items():
        registrations = _parse_integer(_cell_value(row, column_index))
        if registrations is None:
            continue
        records.append(
            RegistrationRecord(
                source_id=KBA_SOURCE_ID,
                year=year,
                month=month,
                brand=brand,
                model_series=model_series,
                registrations=registrations,
                fuel_type=fuel_type,
                market_share=_parse_decimal_percent(
                    _cell_value(row, header_mapping.columns.get("market_share"))
                ),
                source_url=source_url_from_row or source_url,
                collected_at=collected_at,
            )
        )

    return records


def _resolve_period(
    rows: Sequence[Sequence[object]],
    workbook_path: Path,
    sheet_title: str,
    year: int | None,
    month: int | None,
) -> tuple[int, int]:
    if year is not None and month is not None:
        _validate_period(year, month)
        return year, month

    search_parts = [workbook_path.stem, sheet_title]
    for row in rows[:12]:
        search_parts.append(" ".join(str(value) for value in row if value is not None))
    inferred = _infer_period(" ".join(search_parts))

    resolved_year = year if year is not None else inferred[0]
    resolved_month = month if month is not None else inferred[1]
    if resolved_year is None or resolved_month is None:
        raise KBAXlsxParserError(
            f"Unable to determine KBA reporting period for {workbook_path}"
        )
    _validate_period(resolved_year, resolved_month)
    return resolved_year, resolved_month


def _infer_period(text: str) -> tuple[int | None, int | None]:
    normalized = normalize_header(text)
    numeric_match = _PERIOD_NUMERIC_PATTERN.search(normalized)
    if numeric_match is not None:
        return int(numeric_match.group("year")), int(numeric_match.group("month"))

    year_match = _YEAR_PATTERN.search(normalized)
    year = int(year_match.group(1)) if year_match is not None else None
    month = None
    for month_name, month_number in GERMAN_MONTHS.items():
        if re.search(rf"\b{re.escape(month_name)}\b", normalized):
            month = month_number
            break
    return year, month


def _validate_period(year: int, month: int) -> None:
    if not 1900 <= year <= 2100:
        raise KBAXlsxParserError(f"Invalid KBA reporting year: {year}")
    if not 1 <= month <= 12:
        raise KBAXlsxParserError(f"Invalid KBA reporting month: {month}")


def _is_blank_or_footnote_row(row: Sequence[object]) -> bool:
    texts = [_cell_text(row, index) for index in range(len(row))]
    non_empty = [text for text in texts if text]
    if not non_empty:
        return True
    row_text = normalize_header(" ".join(non_empty))
    return any(marker in row_text for marker in FOOTNOTE_MARKERS)


def _should_skip_data_row(brand: str, model_series: str) -> bool:
    if not brand or not model_series:
        return True
    normalized_brand = normalize_header(brand)
    normalized_model = normalize_header(model_series)
    return (
        normalized_brand in SKIP_ROW_MARKERS
        or normalized_model in SKIP_ROW_MARKERS
        or normalized_model.startswith("darunter")
    )


def _cell_value(row: Sequence[object], index: int | None) -> object:
    if index is None or index >= len(row):
        return None
    return row[index]


def _cell_text(row: Sequence[object], index: int | None) -> str:
    value = _cell_value(row, index)
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def _parse_integer(value: object) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, int) and not isinstance(value, bool):
        return value if value >= 0 else None
    if isinstance(value, float):
        return int(value) if value >= 0 and value.is_integer() else None

    match = _NUMBER_PATTERN.search(str(value))
    if match is None:
        return None
    digits = re.sub(r"\D", "", match.group(0))
    if not digits:
        return None
    parsed = int(digits)
    return parsed if parsed >= 0 else None


def _parse_decimal_percent(value: object) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, int | float) and not isinstance(value, bool):
        return Decimal(str(value))

    match = _NUMBER_PATTERN.search(str(value))
    if match is None:
        return None
    normalized = _normalize_decimal_text(match.group(0))
    if not normalized:
        return None
    try:
        return Decimal(normalized)
    except InvalidOperation:
        return None


def _parse_year(value: object) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, int) and not isinstance(value, bool):
        return value if 1900 <= value <= 2100 else None
    match = _YEAR_PATTERN.search(str(value))
    return int(match.group(1)) if match is not None else None


def _parse_month(value: object) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, int) and not isinstance(value, bool):
        return value if 1 <= value <= 12 else None
    text = normalize_header(value)
    if text.isdigit():
        parsed = int(text)
        return parsed if 1 <= parsed <= 12 else None
    return GERMAN_MONTHS.get(text)


def _normalize_decimal_text(value: str) -> str:
    normalized = value.replace(" ", "").strip(".,")
    if "," in normalized and "." in normalized:
        return normalized.replace(".", "").replace(",", ".").strip(".")
    if "," in normalized:
        before, _, after = normalized.partition(",")
        if 1 <= len(after) <= 2:
            return f"{before}.{after}"
        return normalized.replace(",", "")
    if normalized.count(".") == 1:
        before, _, after = normalized.partition(".")
        if 1 <= len(after) <= 2:
            return f"{before}.{after}"
    return normalized.replace(".", "")


def _resolve_collected_at(value: datetime | None) -> datetime:
    if value is None:
        return datetime.now(UTC)
    if value.tzinfo is None:
        return value.replace(tzinfo=UTC)
    return value
