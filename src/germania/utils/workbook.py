"""Shared openpyxl helpers for Project Germania tabular reports."""

from __future__ import annotations

from collections.abc import Iterable, Sequence
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

_HEADER_FILL = PatternFill("solid", fgColor="0F766E")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TABLE_STYLE = "TableStyleMedium2"


def create_report_workbook() -> Workbook:
    """Return a workbook without the default empty worksheet."""

    workbook = Workbook()
    workbook.remove(workbook.active)
    return workbook


def add_table_sheet(
    workbook: Workbook,
    *,
    title: str,
    headers: Sequence[str],
    rows: Iterable[Sequence[Any]],
    table_name: str,
) -> Worksheet:
    """Add one consistently formatted, filterable worksheet."""

    worksheet = workbook.create_sheet(title)
    worksheet.sheet_view.showGridLines = False
    worksheet.append(list(headers))
    for row in rows:
        worksheet.append([_excel_value(value) for value in row])

    for cell in worksheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[1].height = 24
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    if worksheet.max_row > 1:
        table = Table(displayName=table_name, ref=worksheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name=_TABLE_STYLE,
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)

    _format_columns(worksheet, headers)
    return worksheet


def save_report_workbook(workbook: Workbook, path: Path) -> None:
    """Save a workbook after ensuring its output directory exists."""

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _format_columns(worksheet: Worksheet, headers: Sequence[str]) -> None:
    for index, header in enumerate(headers, start=1):
        letter = get_column_letter(index)
        values = [
            worksheet.cell(row=row, column=index).value
            for row in range(1, worksheet.max_row + 1)
        ]
        width = min(
            max(len(str(value)) if value is not None else 0 for value in values) + 2, 48
        )
        worksheet.column_dimensions[letter].width = max(width, 11)

        normalized = header.casefold()
        if "price" in normalized or "amount" in normalized:
            for cell in worksheet[letter][1:]:
                cell.number_format = "#,##0.00"
        elif normalized.endswith("_rate"):
            for cell in worksheet[letter][1:]:
                cell.number_format = "0.0%"
        elif normalized.endswith("_at") or normalized.endswith("_time"):
            for cell in worksheet[letter][1:]:
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
        elif normalized.endswith("_date"):
            for cell in worksheet[letter][1:]:
                cell.number_format = "yyyy-mm-dd"


def _excel_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.replace(tzinfo=None)
    if isinstance(value, date | datetime | str | int | float | bool) or value is None:
        return value
    return str(value)
